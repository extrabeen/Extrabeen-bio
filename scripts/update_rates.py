#!/usr/bin/env python3
"""
Automated rate updater for ExtraBeen Bio & Ads Rates.
Fetches real-time USDT price from Nobitex API and updates ads-rates.html and docx.
"""

import math
import os
import re
import urllib.request
import json
from datetime import datetime, timezone, timedelta

# Fixed base USD prices
BASE_RATES = {
    "story_banner": 37,
    "story_banner_extra": 6,
    "story_selfie": 48,
    "story_selfie_extra": 9,
    "story_banner_selfie": 58,
    "story_banner_selfie_extra": 9,
    "story_serial": 110,
    "story_inperson": 238,
    "post_gadget": 100,
    "post_service": 119,
    "post_event": 311,
    "post_collab": 21,
    "youtube_gadget": 89,
    "youtube_service": 73,
    "package_1": 171,
    "package_2": 161,
}

def gregorian_to_jalali(gy, gm, gd):
    g_d_m = [0, 31, 59, 90, 120, 151, 181, 212, 243, 273, 304, 334]
    if gy > 1600:
        jy = 979
        gy -= 1600
    else:
        jy = 0
        gy -= 621
    gy2 = (gm > 2) and (gy + 1) or gy
    days = (365 * gy) + ((gy2 + 3) // 4) - ((gy2 + 99) // 100) + ((gy2 + 399) // 400) - 80 + gd + g_d_m[gm - 1]
    jy += 33 * (days // 12053)
    days %= 12053
    jy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        jy += (days - 1) // 365
        days = (days - 1) % 365
    if days < 186:
        jm = 1 + (days // 31)
        jd = 1 + (days % 31)
    else:
        jm = 7 + ((days - 186) // 30)
        jd = 1 + ((days - 186) % 30)
    return jy, jm, jd

def get_jalali_date_str():
    # Tehran Timezone: UTC+3:30
    tehran_tz = timezone(timedelta(hours=3, minutes=30))
    now = datetime.now(tehran_tz)
    jy, jm, jd = gregorian_to_jalali(now.year, now.month, now.day)
    months = ['فروردین', 'اردیبهشت', 'خرداد', 'تیر', 'مرداد', 'شهریور', 'مهر', 'آبان', 'آذر', 'دی', 'بهمن', 'اسفند']
    time_str = now.strftime('%H:%M')
    return f"{jd} {months[jm - 1]} {jy}", time_str

def fetch_nobitex_rate():
    urls = [
        "https://apiv2.nobitex.ir/market/stats?srcCurrency=usdt&dstCurrency=rls",
        "https://apiv2.nobitex.ir/market/stats?srcCurrency=usdt&dstCurrency=irt",
    ]
    for url in urls:
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
            with urllib.request.urlopen(req, timeout=10) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                stats = data.get("stats", {})
                for pair in ["usdt-rls", "usdt-irt"]:
                    if pair in stats and "latest" in stats[pair]:
                        rial = float(stats[pair]["latest"])
                        toman = rial / 10 if rial > 1000000 else rial
                        if toman > 50000:
                            return int(toman)
        except Exception as e:
            print(f"Warning: could not fetch from {url}: {e}")
    
    print("Fallback to default/last known rate: 271,380 Tomans")
    return 271380

def format_toman(usd, rate):
    raw_tomans = usd * rate
    t_million = math.ceil(raw_tomans / 100000) / 10
    if t_million == int(t_million):
        s = f"{int(t_million)} میلیون تومن"
    else:
        s = f"{t_million:.1f} میلیون تومن"
    return to_persian_digits(s)

def to_persian_digits(s):
    persian_digits = {'0': '۰', '1': '۱', '2': '۲', '3': '۳', '4': '۴', '5': '۵', '6': '۶', '7': '۷', '8': '۸', '9': '۹'}
    return ''.join(persian_digits.get(c, c) for c in str(s))

def generate_rates(rate):
    calculated = {}
    for key, usd in BASE_RATES.items():
        toman_str = format_toman(usd, rate)
        calculated[key] = {
            "usd": usd,
            "toman_str": toman_str,
            "display": f"{toman_str} (${usd})",
        }
    return calculated

def update_html(calculated, rate, jalali_date, jalali_time, file_path="ads-rates.html"):
    if not os.path.exists(file_path):
        print(f"File {file_path} not found.")
        return False
    
    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Story section replacements
    sb = calculated["story_banner"]["display"]
    sb_extra = f"{calculated['story_banner_extra']['toman_str']} [${calculated['story_banner_extra']['usd']}]"
    content = re.sub(
        r'(<h3 class="text-lg font-bold">تک استوری بنر :</h3>\s*<span class="text-sm font-bold opacity-80">\(هر استوری بیشتر , )[^)]+(\)</span>\s*</div>\s*<span class="text-2xl font-black bg-\[#1f271b\] text-\[#99b19c\] px-3 py-1 self-start md:self-center">)[^<]+(</span>)',
        rf'\g<1>{sb_extra} اضافه میشود\g<2>{sb}\g<3>',
        content
    )

    ss = calculated["story_selfie"]["display"]
    ss_extra = f"{calculated['story_selfie_extra']['toman_str']} [${calculated['story_selfie_extra']['usd']}]"
    content = re.sub(
        r'(<h3 class="text-lg font-bold">تک استوری سلفی انباکس/ معرفی خدمات / اپ :</h3>\s*<span class="text-sm font-bold opacity-80">\(هر استوری بیشتر , )[^)]+(\)</span>\s*</div>\s*<span class="text-2xl font-black bg-\[#1f271b\] text-\[#99b19c\] px-3 py-1 self-start md:self-center">)[^<]+(</span>)',
        rf'\g<1>{ss_extra} اضافه میشود\g<2>{ss}\g<3>',
        content
    )

    sbs = calculated["story_banner_selfie"]["display"]
    sbs_extra = f"{calculated['story_banner_selfie_extra']['toman_str']} [${calculated['story_banner_selfie_extra']['usd']}]"
    content = re.sub(
        r'(<h3 class="text-lg font-bold leading-tight">تک استوری بنر \+ تک استوری سلفی انباکس/ معرفی خدمات / اپ :</h3>\s*<span class="text-sm font-bold opacity-80">\(هر استوری بیشتر , )[^)]+(\)</span>\s*</div>\s*<span class="text-2xl font-black bg-\[#1f271b\] text-\[#99b19c\] px-3 py-1 self-start md:self-center">)[^<]+(</span>)',
        rf'\g<1>{sbs_extra} اضافه میشود\g<2>{sbs}\g<3>',
        content
    )

    # Serial story
    content = re.sub(
        r'(رشته استوری سریالی انباکس/ معرفی خدمات / اپ \( همراه با سناریو \+ فیلمبرداری و ادیت حرفه ای \) :</span>\s*<strong class="whitespace-nowrap text-xl">)[^<]+(</strong>)',
        rf'\g<1>{calculated["story_serial"]["display"]}\g<2>',
        content
    )

    # In-person story
    content = re.sub(
        r'(رشته استوری حضوری \( همراه با سناریو \+ افیش فیلم بردار و ادیت حرفه ای \) :</span>\s*<strong class="whitespace-nowrap text-xl">)[^<]+(</strong>)',
        rf'\g<1>{calculated["story_inperson"]["display"]}\g<2>',
        content
    )

    # Instagram posts
    content = re.sub(
        r'(پست معرفی گجت و آنباکس \( با ادیت حرفه ای و سناریو \):</span>\s*<strong class="whitespace-nowrap text-xl bg-\[#1f271b\] text-\[#99b19c\] px-2 self-start md:self-center">)[^<]+(</strong>)',
        rf'\g<1>{calculated["post_gadget"]["display"]}\g<2>',
        content
    )

    content = re.sub(
        r'(پست معرفی خدمات \( سایت / اپ \) \( با ادیت حرفه ای و سناریو \) :</span>\s*<strong class="whitespace-nowrap text-xl bg-\[#1f271b\] text-\[#99b19c\] px-2 self-start md:self-center">)[^<]+(</strong>)',
        rf'\g<1>{calculated["post_service"]["display"]}\g<2>',
        content
    )

    content = re.sub(
        r'(پست رویداد و معرفی حضوری \( همراه با سناریو \+ آفیش فیلمبردار و ادیت حرفه ای \) :</span>\s*<strong class="whitespace-nowrap text-xl bg-\[#1f271b\] text-\[#99b19c\] px-2 self-start md:self-center">)[^<]+(</strong>)',
        rf'\g<1>{calculated["post_event"]["display"]}\g<2>',
        content
    )

    collab_disp = calculated["post_collab"]["display"]
    content = re.sub(
        r'(پست Collaboration با برند : به تعرفه های بالا مبلغ )[^<]+( اضافه می شود\.)',
        rf'\g<1>{collab_disp}\g<2>',
        content
    )

    # YouTube
    content = re.sub(
        r'(پست معرفی گجت و آنباکس \( با ادیت حرفه ای \+ کپشن سئو شده \):</span>\s*<strong class="whitespace-nowrap text-xl">)[^<]+(</strong>)',
        rf'\g<1>{calculated["youtube_gadget"]["display"]}\g<2>',
        content
    )

    content = re.sub(
        r'(پست معرفی خدمات \( سایت / اپ \) \(با ادیت حرفه ای \+ کپشن سئو شده \) :</span>\s*<strong class="whitespace-nowrap text-xl">)[^<]+(</strong>)',
        rf'\g<1>{calculated["youtube_service"]["display"]}\g<2>',
        content
    )

    # Packages
    content = re.sub(
        r'(پکیج ۱</span>\s*پست معرفی گجت و آنباکس[^<]+</p>\s*<div class="lg:self-center[^>]+>\s*<span class="text-2xl font-black whitespace-nowrap bg-\[#1f271b\] text-\[#99b19c\] px-3 py-1 block text-center">)[^<]+(</span>)',
        rf'\g<1>{calculated["package_1"]["display"]}\g<2>',
        content
    )

    content = re.sub(
        r'(پکیج ۲</span>\s*پست معرفی خدمات سایت / اپ[^<]+</p>\s*<div class="lg:self-center[^>]+>\s*<span class="text-2xl font-black whitespace-nowrap bg-\[#1f271b\] text-\[#99b19c\] px-3 py-1 block text-center">)[^<]+(</span>)',
        rf'\g<1>{calculated["package_2"]["display"]}\g<2>',
        content
    )

    # Nobitex rate info & auto-update timestamp
    rate_formatted = f"{rate:,}"
    rate_persian = to_persian_digits(rate_formatted)
    date_persian = to_persian_digits(jalali_date)
    time_persian = to_persian_digits(jalali_time)

    # Replace or add rate info in notes
    rate_note = f'نرخ لحظه‌ای تتر در نوبیتکس : {rate_persian} تومان'
    content = re.sub(
        r'(<span>نرخ مبنای محاسبه هر دلار : )[^<]+(</span>)',
        rf'<span>{rate_note}</span>',
        content
    )
    content = re.sub(
        r'(<span>نرخ لحظه‌ای تتر در نوبیتکس : )[^<]+(</span>)',
        rf'<span>{rate_note}</span>',
        content
    )

    # Replace date box
    date_text = f'آخرین به‌روزرسانی خودکار : {date_persian} - ساعت {time_persian}'
    content = re.sub(
        r'(<div class="mt-6 pt-4 border-t-2 border-\[#99b19c\] border-dashed text-sm font-mono font-bold text-center">\s*)[^<]+(\s*</div>)',
        rf'\g<1>{date_text}\g<2>',
        content
    )

    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content)

    print(f"Updated {file_path} successfully with USDT rate = {rate_formatted} Tomans.")
    return True

def update_docx(calculated, rate, jalali_date, file_path="تعرفه های اکسترابین.docx"):
    if not os.path.exists(file_path):
        return False
    try:
        import docx
        doc = docx.Document(file_path)
        rate_persian = to_persian_digits(f"{rate:,}")
        date_persian = to_persian_digits(jalali_date)

        # Exact paragraph assignments to avoid any collision
        sb_extra = f"{calculated['story_banner_extra']['toman_str']} [${calculated['story_banner_extra']['usd']}]"
        doc.paragraphs[3].text = f"تک استوری بنر : {calculated['story_banner']['display']} (هر استوری بیشتر , {sb_extra} اضافه میشود)"

        ss_extra = f"{calculated['story_selfie_extra']['toman_str']} [${calculated['story_selfie_extra']['usd']}]"
        doc.paragraphs[4].text = f"تک استوری سلفی انباکس/ معرفی خدمات / اپ  :  {calculated['story_selfie']['display']} (هر استوری بیشتر , {ss_extra} اضافه میشود)"

        sbs_extra = f"{calculated['story_banner_selfie_extra']['toman_str']} [${calculated['story_banner_selfie_extra']['usd']}]"
        doc.paragraphs[5].text = f"تک استوری بنر + تک استوری سلفی انباکس/ معرفی خدمات / اپ : {calculated['story_banner_selfie']['display']} (هر استوری بیشتر , {sbs_extra} اضافه میشود)"

        doc.paragraphs[7].text = f"رشته استوری سریالی انباکس/ معرفی خدمات / اپ ( همراه با سناریو + فیلمبرداری و ادیت حرفه ای ) : {calculated['story_serial']['display']}"
        doc.paragraphs[9].text = f"رشته استوری حضوری ( همراه با سناریو + افیش فیلم بردار و ادیت حرفه ای ) : {calculated['story_inperson']['display']}"

        doc.paragraphs[15].text = f"پست معرفی گجت و آنباکس ( با ادیت حرفه ای و سناریو ): {calculated['post_gadget']['display']}"
        doc.paragraphs[16].text = f"پست معرفی خدمات ( سایت / اپ ) (با ادیت حرفه ای و سناریو ) : {calculated['post_service']['display']}"
        doc.paragraphs[18].text = f"پست رویداد و معرفی حضوری ( همراه با سناریو + آفیش فیلمبردار و ادیت حرفه ای ) : {calculated['post_event']['display']}"
        doc.paragraphs[20].text = f"پست  Collaborationبا برند : به تعرفه های بالا مبلغ {calculated['post_collab']['display']} اضافه می شود."

        doc.paragraphs[32].text = f"پست معرفی گجت و آنباکس ( با ادیت حرفه ای + کپشن سئو شده ): {calculated['youtube_gadget']['display']}"
        doc.paragraphs[33].text = f"پست معرفی خدمات ( سایت / اپ ) (با ادیت حرفه ای + کپشن سئو شده ) : {calculated['youtube_service']['display']}"

        doc.paragraphs[49].text = f"پست معرفی گجت و آنباکس ( با ادیت حرفه ای )  در یوتیوب / آپارات ( ویدیو بلند  +  کپشن سئو شده ) + ویدیو کوتاه مدت اینستاگرام / شورت یوتیوب و آپارات ( با ادیت حرفه ای + سناریو )  : {calculated['package_1']['display']}"
        doc.paragraphs[50].text = f"\nپست معرفی خدمات سایت / اپ ( با ادیت حرفه ای )  در یوتیوب / آپارات ( ویدیو بلند  +  کپشن سئو شده ) + ویدیو کوتاه مدت اینستاگرام / شورت یوتیوب و آپارات ( با ادیت حرفه ای + سناریو )  : {calculated['package_2']['display']}"

        doc.paragraphs[56].text = f"تاریخ تنظیم تعرفه : {date_persian} (نرخ لحظه‌ای تتر نوبیتکس: {rate_persian} تومان)"

        doc.save(file_path)
        print(f"Updated {file_path} successfully.")
        return True
    except Exception as e:
        print(f"Docx update skipped or failed: {e}")
        return False

def main():
    rate = fetch_nobitex_rate()
    jalali_date, jalali_time = get_jalali_date_str()
    print(f"Current Nobitex USDT Rate: {rate:,} Tomans ({jalali_date} {jalali_time})")

    calculated = generate_rates(rate)
    update_html(calculated, rate, jalali_date, jalali_time)
    update_docx(calculated, rate, jalali_date)

if __name__ == "__main__":
    main()
